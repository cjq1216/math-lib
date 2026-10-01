"""试卷文档（Word、PDF、纯文本/Markdown）多格式文本提取引擎。

支持格式：
1. Word (.docx)：解析正文段落与表格文本；
2. PDF (.pdf)：基于 pypdf 提取矢量/电子排版文字与页码；对纯图片扫描件自动拦截并返回引导提示；
3. 纯文本 / Markdown (.txt, .md)：UTF-8 / GBK 编码自适应解码。
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

import docx
import pypdf


class DocumentExtractionError(Exception):
    """文档解析异常"""

    def __init__(self, message: str, code: str = "extraction_failed"):
        super().__init__(message)
        self.message = message
        self.code = code


@dataclass
class ExtractedDocument:
    """提取后的文档内容结构"""

    filename: str
    file_type: str  # "docx" | "pdf" | "text"
    total_chars: int
    page_count: int | None
    content: str


def extract_document_text(filename: str, file_bytes: bytes) -> ExtractedDocument:
    """从二进制数据流中解析出规范化试卷文本。"""
    if not file_bytes:
        raise DocumentExtractionError("上传的文件内容为空", code="empty_file")

    ext = Path(filename).suffix.lower()

    if ext == ".docx":
        return _extract_docx(filename, file_bytes)
    elif ext == ".pdf":
        return _extract_pdf(filename, file_bytes)
    elif ext in (".txt", ".md", ".text"):
        return _extract_text(filename, file_bytes)
    else:
        raise DocumentExtractionError(
            f"不支持的文件格式 '{ext}'。目前支持 Word (.docx)、PDF (.pdf) 与纯文本 (.txt, .md)",
            code="unsupported_format",
        )


def _extract_docx(filename: str, file_bytes: bytes) -> ExtractedDocument:
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise DocumentExtractionError(f"Word 文档解析失败: {e}", code="docx_parse_error") from e

    lines: list[str] = []

    # 1. 提取正文段落
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            lines.append(txt)

    # 2. 提取表格内容
    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells]
            if any(row_texts):
                lines.append(" | ".join(row_texts))

    content = "\n\n".join(lines).strip()
    if not content:
        raise DocumentExtractionError("Word 文档未包含任何有效文字内容", code="empty_content")

    return ExtractedDocument(
        filename=filename,
        file_type="docx",
        total_chars=len(content),
        page_count=None,
        content=content,
    )


def _extract_pdf(filename: str, file_bytes: bytes) -> ExtractedDocument:
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    except Exception as e:
        raise DocumentExtractionError(f"PDF 文档解析失败: {e}", code="pdf_parse_error") from e

    page_count = len(reader.pages)
    if page_count == 0:
        raise DocumentExtractionError("PDF 文档无有效页面", code="empty_pdf")

    page_texts: list[str] = []
    total_valid_chars = 0

    for page_idx, page in enumerate(reader.pages, start=1):
        try:
            raw_text = page.extract_text() or ""
        except Exception:
            raw_text = ""
        clean_text = raw_text.strip()
        if clean_text:
            total_valid_chars += len(clean_text)
            page_texts.append(f"--- 第 {page_idx} 页 ---\n{clean_text}")

    # 扫描版 PDF 判定：如果所有页面提取文字累计过少（少于 15 字符），则视为纯扫描图片版
    if total_valid_chars < 15:
        raise DocumentExtractionError(
            "检测到该 PDF 疑似为纯扫描图片版或无文字排版。系统支持电子文字版 PDF、Word (.docx) 与文本文件；"
            "建议使用微信、WPS 或在线 OCR 工具将图片识别转换为 Word 文档后再上传切题。",
            code="scanned_image_pdf",
        )

    content = "\n\n".join(page_texts).strip()
    return ExtractedDocument(
        filename=filename,
        file_type="pdf",
        total_chars=len(content),
        page_count=page_count,
        content=content,
    )


def _extract_text(filename: str, file_bytes: bytes) -> ExtractedDocument:
    for encoding in ("utf-8", "gb18030", "gbk"):
        try:
            content = file_bytes.decode(encoding).strip()
            if not content:
                raise DocumentExtractionError("文本文件内容为空", code="empty_content")
            return ExtractedDocument(
                filename=filename,
                file_type="text",
                total_chars=len(content),
                page_count=1,
                content=content,
            )
        except UnicodeDecodeError:
            continue

    raise DocumentExtractionError("文本文件编码无法识别，请使用 UTF-8 编码保存", code="encoding_error")
