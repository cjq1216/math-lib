"""学生录入、查询和 Excel 批量导入。"""

import io
from datetime import datetime
from typing import Annotated

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import or_
from sqlmodel import Session, select

from app.core.database import get_session
from app.core.dependencies import (
    can_access_class,
    can_access_student,
    get_current_user,
    require_student_access,
)
from app.models.class_ import Class, ClassStudent, ClassTeacher
from app.models.student import Student
from app.models.user import User, UserRole
from app.schemas.common import ImportErrorItem, ImportSummary, OkResponse
from app.schemas.student import StudentCreate, StudentRead, StudentUpdate
from app.services.audit_service import add_audit_event

router = APIRouter()


def _student_to_dict(student: Student) -> dict:
    return {
        "id": student.id,
        "student_no": student.student_no,
        "name": student.name,
        "gender": student.gender,
        "grade": student.grade,
        "enrollment_year": student.enrollment_year,
        "phone": student.phone,
        "parent_phone": student.parent_phone,
        "average_score": student.average_score,
        "notes": student.notes,
        "is_active": student.is_active,
        "created_at": student.created_at,
    }


def _join_student_to_class(session: Session, student_id: int, class_id: int) -> None:
    relation = session.exec(
        select(ClassStudent).where(
            ClassStudent.class_id == class_id,
            ClassStudent.student_id == student_id,
        )
    ).first()
    if relation is None:
        session.add(ClassStudent(class_id=class_id, student_id=student_id))
    else:
        relation.left_at = None
        relation.joined_at = datetime.utcnow()
        session.add(relation)


@router.post("/", response_model=StudentRead, status_code=status.HTTP_201_CREATED)
def create_student(
    payload: StudentCreate,
    request: Request,
    session: Annotated[Session, Depends(get_session)],
    current_user: Annotated[User, Depends(get_current_user)],
) -> dict:
    """创建学生，可同时加入一个有权访问的班级。"""
    if session.exec(
        select(Student.id).where(Student.student_no == payload.student_no)
    ).first() is not None:
        raise HTTPException(status_code=400, detail="学号已存在")
    if payload.class_id is not None and not can_access_class(
        session, current_user, payload.class_id
    ):
        raise HTTPException(status_code=403, detail="无权向该班级添加学生")

    values = payload.model_dump(exclude={"class_id"})
    student = Student(**values, created_by=current_user.id)
    session.add(student)
    session.flush()
    if payload.class_id is not None:
        _join_student_to_class(session, student.id, payload.class_id)
    add_audit_event(
        session,
        action="create",
        resource_type="student",
        actor=current_user,
        resource_id=student.id,
        changes={"student_no": student.student_no, "class_id": payload.class_id},
        request=request,
    )
    session.commit()
    return _student_to_dict(student)


@router.get("/", response_model=list[StudentRead])
def list_students(
    session: Annotated[Session, Depends(get_session)],
    current_user: Annotated[User, Depends(get_current_user)],
    grade: int | None = Query(default=None, ge=7, le=9),
    is_active: bool = True,
    skip: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=100),
    keyword: str | None = None,
) -> list[dict]:
    """管理员查看全部；教师仅查看任教班级的在班学生。"""
    statement = select(Student).where(Student.is_active == is_active)
    if current_user.role == UserRole.TEACHER:
        statement = (
            statement.join(ClassStudent, ClassStudent.student_id == Student.id)
            .join(Class, Class.id == ClassStudent.class_id)
            .outerjoin(ClassTeacher, ClassTeacher.class_id == Class.id)
            .where(ClassStudent.left_at.is_(None))
            .where(
                or_(
                    Class.head_teacher_id == current_user.id,
                    ClassTeacher.teacher_id == current_user.id,
                )
            )
            .distinct()
        )
    if grade:
        statement = statement.where(Student.grade == grade)
    if keyword:
        statement = statement.where(
            Student.name.contains(keyword) | Student.student_no.contains(keyword)
        )
    students = session.exec(
        statement.order_by(Student.grade, Student.student_no).offset(skip).limit(limit)
    ).all()
    return [_student_to_dict(student) for student in students]


@router.get("/template", response_class=StreamingResponse)
def download_student_template() -> StreamingResponse:
    """下载学生导入模板。"""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "学生信息"
    headers = ["学号", "姓名", "性别", "年级", "联系电话", "家长电话", "备注"]
    worksheet.append(headers)
    worksheet.append(["2024001", "张三", "男", 7, "138xxxx", "139xxxx", ""])
    worksheet.append(["2024002", "李四", "女", 7, "138xxxx", "", ""])
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=students_template.xlsx"},
    )


@router.get("/{student_id}", response_model=StudentRead)
def get_student(
    student: Annotated[Student, Depends(require_student_access)],
) -> dict:
    """读取有权访问的学生。"""
    return _student_to_dict(student)


@router.patch("/{student_id}", response_model=StudentRead)
def update_student(
    payload: StudentUpdate,
    request: Request,
    student: Annotated[Student, Depends(require_student_access)],
    session: Annotated[Session, Depends(get_session)],
    current_user: Annotated[User, Depends(get_current_user)],
) -> dict:
    """按白名单更新学生。"""
    updates = payload.model_dump(exclude_unset=True)
    for field, value in updates.items():
        setattr(student, field, value)
    student.updated_at = datetime.utcnow()
    session.add(student)
    add_audit_event(
        session,
        action="update",
        resource_type="student",
        actor=current_user,
        resource_id=student.id,
        changes={"fields": sorted(updates)},
        request=request,
    )
    session.commit()
    return _student_to_dict(student)


@router.post("/import", response_model=ImportSummary)
async def import_students(
    request: Request,
    session: Annotated[Session, Depends(get_session)],
    current_user: Annotated[User, Depends(get_current_user)],
    file: UploadFile = File(...),
    class_id: int | None = None,
) -> ImportSummary:
    """Excel 批量导入学生，错误返回稳定的行列结构。"""
    if class_id is not None and not can_access_class(session, current_user, class_id):
        raise HTTPException(status_code=403, detail="无权向该班级导入学生")

    content = await file.read()
    try:
        workbook = load_workbook(io.BytesIO(content))
    except Exception as error:
        raise HTTPException(status_code=400, detail=f"无法读取 Excel: {error}") from error
    worksheet = workbook.active
    headers = [cell.value for cell in worksheet[1]]
    expected = ["学号", "姓名", "性别", "年级", "联系电话", "家长电话", "备注"]
    if not all(header in headers for header in ["学号", "姓名", "年级"]):
        raise HTTPException(status_code=400, detail=f"Excel 至少需要包含列：{expected}")

    columns = {header: index for index, header in enumerate(headers) if header}
    created = 0
    updated = 0
    errors: list[ImportErrorItem] = []
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        if not row[columns["学号"]]:
            continue
        try:
            student_no = str(row[columns["学号"]]).strip()
            name = str(row[columns["姓名"]]).strip()
            grade = int(row[columns["年级"]])
            if grade not in {7, 8, 9}:
                raise ValueError("年级必须是 7、8 或 9")

            student = session.exec(
                select(Student).where(Student.student_no == student_no)
            ).first()
            if student is not None:
                if (
                    current_user.role != UserRole.ADMIN
                    and not can_access_student(session, current_user, student.id)
                    and student.created_by != current_user.id
                ):
                    errors.append(
                        ImportErrorItem(
                            row=row_index,
                            column="学号",
                            code="student_forbidden",
                            message=f"无权更新学号 {student_no}",
                        )
                    )
                    continue
                student.name = name
                student.grade = grade
                if "性别" in columns and row[columns["性别"]]:
                    student.gender = str(row[columns["性别"]])
                student.updated_at = datetime.utcnow()
                session.add(student)
                updated += 1
            else:
                student = Student(
                    student_no=student_no,
                    name=name,
                    gender=(str(row[columns["性别"]]) if "性别" in columns and row[columns["性别"]] else None),
                    grade=grade,
                    phone=(str(row[columns["联系电话"]]) if "联系电话" in columns and row[columns["联系电话"]] else None),
                    parent_phone=(str(row[columns["家长电话"]]) if "家长电话" in columns and row[columns["家长电话"]] else None),
                    notes=(str(row[columns["备注"]]) if "备注" in columns and row[columns["备注"]] else None),
                    created_by=current_user.id,
                )
                session.add(student)
                session.flush()
                created += 1
            if class_id is not None:
                _join_student_to_class(session, student.id, class_id)
        except Exception as error:
            errors.append(
                ImportErrorItem(
                    row=row_index,
                    code="invalid_row",
                    message=str(error),
                )
            )

    add_audit_event(
        session,
        action="import",
        resource_type="student",
        actor=current_user,
        changes={
            "class_id": class_id,
            "created": created,
            "updated": updated,
            "error_count": len(errors),
        },
        request=request,
    )
    session.commit()
    return ImportSummary(created=created, updated=updated, errors=errors)


@router.delete("/{student_id}", response_model=OkResponse)
def delete_student(
    request: Request,
    student: Annotated[Student, Depends(require_student_access)],
    session: Annotated[Session, Depends(get_session)],
    current_user: Annotated[User, Depends(get_current_user)],
) -> OkResponse:
    """软删有权访问的学生。"""
    student.is_active = False
    student.updated_at = datetime.utcnow()
    session.add(student)
    add_audit_event(
        session,
        action="delete",
        resource_type="student",
        actor=current_user,
        resource_id=student.id,
        request=request,
    )
    session.commit()
    return OkResponse()
