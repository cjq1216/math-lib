"""R5 线上 LLM 任务、文档切题入库、向量相似题检索与平滑降级集成测试。"""

import io

import docx
import pypdf
from fastapi.testclient import TestClient
from sqlmodel import Session

from app.models.background_task import TaskStatus, TaskType
from app.models.question import Question, QuestionType
from app.services.llm_client import extract_json_payload
from app.services.llm_service import embed_text
from app.services.task_service import (
    create_task,
    finish_task,
    recover_orphaned_tasks,
    start_task,
    update_progress,
)
from app.services.vector_service import (
    cosine_similarity,
    save_question_embedding,
)


def _init_auth(client: TestClient) -> tuple[dict[str, str], dict[str, str]]:
    """初始化管理员与测试教师。"""
    res = client.post(
        "/api/v1/auth/register",
        json={"username": "admin", "password": "password123", "real_name": "系统管理员"},
    )
    assert res.status_code == 200
    admin_token = res.json()["access_token"]
    admin_headers = {"Authorization": f"Bearer {admin_token}"}

    res = client.post(
        "/api/v1/users/",
        headers=admin_headers,
        json={"username": "teacher1", "password": "password123", "real_name": "数学张老师", "role": "teacher"},
    )
    assert res.status_code == 201

    login_res = client.post("/api/v1/auth/login", data={"username": "teacher1", "password": "password123"})
    assert login_res.status_code == 200
    teacher_token = login_res.json()["access_token"]
    teacher_headers = {"Authorization": f"Bearer {teacher_token}"}

    return admin_headers, teacher_headers


def _create_mock_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("2026年秋季期末考试", level=1)
    doc.add_paragraph("1. 已知一元二次方程 $x^2 - 4 = 0$，求方程的解。")
    doc.add_paragraph("A. 2   B. -2   C. ±2   D. 4")
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def _create_empty_pdf() -> bytes:
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=595, height=842)
    bio = io.BytesIO()
    writer.write(bio)
    return bio.getvalue()


def test_split_document_upload_and_task_creation(client: TestClient):
    """测试通过上传试卷文档创建后台切题任务及扫描件拦截。"""
    _, teacher_h = _init_auth(client)

    # 1. 上传正常 Word 文档
    docx_bytes = _create_mock_docx()
    res = client.post(
        "/api/v1/llm/split-doc",
        headers=teacher_h,
        files={"file": ("mock_exam.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res.status_code == 200
    data = res.json()
    assert data["task_id"] is not None
    assert data["status"] == "pending"
    assert "文档解析成功" in data["message"]

    # 查询任务详情
    task_res = client.get(f"/api/v1/llm/tasks/{data['task_id']}", headers=teacher_h)
    assert task_res.status_code == 200
    task_info = task_res.json()
    assert task_info["task_type"] == "split_exam"

    # 2. 上传无文字的扫描图片版 PDF：友好拦截
    empty_pdf_bytes = _create_empty_pdf()
    scan_res = client.post(
        "/api/v1/llm/split-doc",
        headers=teacher_h,
        files={"file": ("scan.pdf", io.BytesIO(empty_pdf_bytes), "application/pdf")},
    )
    assert scan_res.status_code == 422
    assert "纯扫描图片版" in scan_res.json()["detail"]


def test_task_service_state_machine_and_recovery(test_engine):
    """测试后台任务状态机流转与应用启动孤儿任务自愈。"""
    with Session(test_engine) as session:
        # 创建任务
        task = create_task(session, TaskType.SPLIT_EXAM, created_by=1)
        assert task.status == TaskStatus.PENDING

        # 开始
        start_task(session, task.id)
        assert task.status == TaskStatus.RUNNING

        # 进度
        update_progress(session, task.id, 50, "正在处理...")
        assert task.progress == 50

        # 完成（部分成功）
        finish_task(session, task.id, result={"questions": [{"number": "1"}]}, partial=True)
        assert task.status == TaskStatus.PARTIAL_SUCCESS
        assert task.progress == 100

        # 创建一个模拟进程中断遗留的 running 任务
        orphan = create_task(session, TaskType.AUTO_TAG, created_by=1)
        start_task(session, orphan.id)
        assert orphan.status == TaskStatus.RUNNING

        # 触发应用启动自愈
        recovered_count = recover_orphaned_tasks(session)
        assert recovered_count >= 1

        session.refresh(orphan)
        assert orphan.status == TaskStatus.FAILED
        assert "服务重启导致任务中断" in (orphan.error_message or "")


def test_batch_curate_commit_drafts_into_question_bank(client: TestClient):
    """测试教师校对草稿后一键批量正式入库。"""
    _, teacher_h = _init_auth(client)

    # 准备知识点
    client.post("/api/v1/knowledge/", headers=teacher_h, json={"code": "K-ALG", "name": "代数综合", "grade": 7})

    drafts = [
        {
            "number": "1",
            "stem": "计算 $(-2)^2$ 的值。",
            "question_type": "choice_single",
            "difficulty": 1,
            "total_score": 3.0,
            "options": ["A. 4", "B. -4", "C. 2", "D. -2"],
            "answer": "A",
            "analysis": "负数的偶次方为正数。",
            "knowledge_points": ["代数综合"],
        },
        {
            "number": "2",
            "stem": "解方程组：(1) $x+y=5$; (2) $x-y=1$。",
            "question_type": "solution",
            "difficulty": 3,
            "total_score": 10.0,
            "answer": "x=3, y=2",
            "analysis": "两式相加即可消去 y。",
            "sub_questions": [
                {"label": "(1)", "stem": "求 x 的值", "score": 5.0, "answer": "x=3"},
                {"label": "(2)", "stem": "求 y 的值", "score": 5.0, "answer": "y=2"},
            ],
            "knowledge_points": [],
        },
    ]

    commit_res = client.post(
        "/api/v1/llm/commit-drafts",
        headers=teacher_h,
        json={"questions": drafts, "source_id": None},
    )
    assert commit_res.status_code == 200
    res_data = commit_res.json()
    assert res_data["created_count"] == 2
    q_ids = res_data["question_ids"]
    assert len(q_ids) == 2

    # 验证题库中题目已完整入库并打上已校对标记
    q1 = client.get(f"/api/v1/questions/{q_ids[0]}", headers=teacher_h).json()
    assert q1["stem"] == "计算 $(-2)^2$ 的值。"
    assert q1["is_verified"] is True
    assert len(q1["options"]) == 4
    assert q1["checksum"] is not None

    q2 = client.get(f"/api/v1/questions/{q_ids[1]}", headers=teacher_h).json()
    assert len(q2["sub_questions"]) == 2
    assert q2["is_verified"] is True


def test_vector_similarity_search_and_graceful_degradation(client: TestClient, test_engine):
    """测试内存余弦相似度检索准确性与无 Key / 故障时的平滑降级。"""
    _, teacher_h = _init_auth(client)

    # 1. 验证余弦相似度数学函数
    assert cosine_similarity([1.0, 0.0], [1.0, 0.0]) == 1.0
    assert cosine_similarity([1.0, 0.0], [0.0, 1.0]) == 0.0
    assert cosine_similarity([], []) == 0.0

    # 2. 在数据库中创建 3 道题并持久化 Embedding
    with Session(test_engine) as session:
        q1 = Question(stem="求一元二次方程根", question_type=QuestionType.CHOICE_SINGLE, difficulty=2, is_active=True)
        q2 = Question(stem="一元二次方程求根公式题", question_type=QuestionType.CHOICE_SINGLE, difficulty=2, is_active=True)
        q3 = Question(stem="全等三角形证明", question_type=QuestionType.PROOF, difficulty=4, is_active=True)
        session.add_all([q1, q2, q3])
        session.commit()
        session.refresh(q1)
        session.refresh(q2)
        session.refresh(q3)

        # q1 与 q2 方向非常接近（相似），q3 正交（不相似）
        save_question_embedding(session, q1.id, [0.99, 0.05, 0.0])
        save_question_embedding(session, q2.id, [0.95, 0.10, 0.0])
        save_question_embedding(session, q3.id, [0.0, 0.0, 1.0])
        session.commit()

        q1_id, q2_id, _ = q1.id, q2.id, q3.id

    # 3. 通过 API 查询 q1 的相似题
    sim_res = client.get(f"/api/v1/llm/questions/{q1_id}/similar?top_k=2&threshold=0.8", headers=teacher_h)
    assert sim_res.status_code == 200
    similar_items = sim_res.json()
    assert len(similar_items) == 1
    # 命中 q2，相似度 > 0.95
    assert similar_items[0]["question_id"] == q2_id
    assert similar_items[0]["similarity"] > 0.95

    # 4. 测试 Embedding 平滑降级（无 Key 时返回 None，绝不抛出 500 异常）
    # 当 MINIMAX_API_KEY 为空时
    import asyncio
    vector_result = asyncio.run(embed_text("测试题干文本"))
    # 在未配置真实有效 key 时返回 None，平滑跳过
    assert vector_result is None or isinstance(vector_result, list)


def test_llm_json_extractor_robustness():
    """测试从 Markdown 代码块与对话杂质中提取 JSON 的鲁棒性。"""
    # 纯 JSON
    assert extract_json_payload('{"key": "value"}') == {"key": "value"}

    # Markdown 代码块包裹
    fenced = """```json
    {
      "questions": [{"number": "1"}]
    }
    ```"""
    assert extract_json_payload(fenced)["questions"][0]["number"] == "1"

    # 首尾带闲聊文字
    messy = """这是我为你切分的试卷题目：
    ```
    [
      {"stem": "题目1"}
    ]
    ```
    希望对你有帮助！"""
    res = extract_json_payload(messy)
    assert isinstance(res, list)
    assert res[0]["stem"] == "题目1"
